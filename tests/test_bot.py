from datetime import date

import pytest

from bot.exportar import a_excel, tablas
from bot.extraer import a_numero, parse_fecha, parse_texto
from bot.validar import problemas

ANDINA = """DISTRIBUIDORA ANDINA S.A.S.
NIT: 900.123.456-7
FACTURA ELECTRÓNICA DE VENTA No. FE-1020
Fecha de emisión: 21/08/2026
Cuaderno argollado 1 $ 12.400 $ 12.400
Resma papel carta 5 $ 18.500 $ 92.500
Subtotal: $ 104.900
IVA 19%: $ 19.931
TOTAL A PAGAR: $ 124.831"""

TORNILLO = """Ferretería El Tornillo
Nit 800987654-1 - Régimen común
Factura N° 4581
Fecha: 2026-08-07
Silicona transparente 2 13.200 26.400
Valor antes de IVA 26.400
IVA 5.016
Total 31.416"""


@pytest.mark.parametrize("texto,esperado", [("1.487.500", 1_487_500), ("1487500,50", 1_487_500.5), ("12.400", 12_400)])
def test_numeros_formato_colombiano(texto, esperado):
    assert a_numero(texto) == esperado


@pytest.mark.parametrize("texto,esperado", [
    ("Fecha de emisión: 21/08/2026", date(2026, 8, 21)),
    ("Fecha: 2026-08-07", date(2026, 8, 7)),
    ("Fecha: 18 de septiembre de 2026", date(2026, 9, 18)),
])
def test_tres_formatos_de_fecha(texto, esperado):
    assert parse_fecha(texto) == esperado


def test_extrae_factura_completa():
    f = parse_texto(ANDINA)
    assert (f.proveedor, f.nit, f.numero) == ("DISTRIBUIDORA ANDINA S.A.S.", "900.123.456-7", "FE-1020")
    assert (f.subtotal, f.iva, f.total) == (104_900, 19_931, 124_831)
    assert len(f.items) == 2 and f.items[1]["cantidad"] == 5


def test_iva_no_se_confunde_con_valor_antes_de_iva():
    f = parse_texto(TORNILLO)
    assert f.subtotal == 26_400 and f.iva == 5_016


def test_factura_correcta_no_tiene_problemas():
    assert problemas(parse_texto(ANDINA), set()) == []


def test_detecta_total_que_no_cuadra():
    f = parse_texto(ANDINA.replace("124.831", "174.831"))
    assert any("no cuadra" in p for p in problemas(f, set()))


def test_detecta_duplicado():
    vistas = set()
    problemas(parse_texto(ANDINA), vistas)
    assert any("duplicada" in p for p in problemas(parse_texto(ANDINA), vistas))


def test_detecta_nit_faltante():
    f = parse_texto(ANDINA.replace("NIT: 900.123.456-7\n", ""))
    assert "Falta NIT del proveedor" in problemas(f, set())


def test_excel_tiene_las_cuatro_hojas(tmp_path):
    from openpyxl import load_workbook
    ok, malo = parse_texto(ANDINA, "a.pdf"), parse_texto(TORNILLO.replace("31.416", "99.999"), "b.pdf")
    vistas = set()
    hojas = tablas([(ok, problemas(ok, vistas)), (malo, problemas(malo, vistas))])
    a_excel(hojas, tmp_path / "x.xlsx")
    wb = load_workbook(tmp_path / "x.xlsx")
    assert wb.sheetnames == ["Resumen", "Facturas", "Detalle", "Revisar"]
    assert hojas["Resumen"].Total.sum() == 124_831       # solo suma las facturas OK
    assert len(hojas["Revisar"]) == 1


def test_detalle_no_repite_productos_de_duplicados():
    vistas = set()
    a, b = parse_texto(ANDINA, "a.pdf"), parse_texto(ANDINA, "a_copia.pdf")
    hojas = tablas([(a, problemas(a, vistas)), (b, problemas(b, vistas))])
    assert len(hojas["Detalle"]) == 2
