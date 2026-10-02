"""Construye el Excel de salida (4 hojas con formato) y el reporte en texto."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = PatternFill("solid", fgColor="0F2A43")
ROJO = PatternFill("solid", fgColor="FDE2E1")
MONEDA = '"$"#,##0'


def tablas(resultados: list[tuple]) -> dict[str, pd.DataFrame]:
    """resultados = [(Factura, [problemas])] -> DataFrames para cada hoja."""
    fact, det, rev = [], [], []
    for f, probs in resultados:
        fila = {"Archivo": f.archivo, "Proveedor": f.proveedor, "NIT": f.nit, "Número": f.numero,
                "Fecha": f.fecha, "Subtotal": f.subtotal, "IVA": f.iva, "Total": f.total,
                "Estado": "OK" if not probs else "REVISAR"}
        fact.append(fila)
        for i in (f.items if not probs else []):     # el detalle solo cuenta facturas aprobadas
            det.append({"Número": f.numero, "Proveedor": f.proveedor, **{k.replace("_", " ").capitalize(): v for k, v in i.items()}})
        for pr in probs:
            rev.append({"Archivo": f.archivo, "Proveedor": f.proveedor, "Número": f.numero, "Problema": pr})
    facturas = pd.DataFrame(fact)
    ok = facturas[facturas.Estado == "OK"].copy()
    ok["Mes"] = pd.to_datetime(ok.Fecha).dt.strftime("%Y-%m")
    resumen = (ok.groupby(["Mes", "Proveedor"], as_index=False)
               .agg(Facturas=("Número", "count"), Subtotal=("Subtotal", "sum"), IVA=("IVA", "sum"), Total=("Total", "sum")))
    return {"Resumen": resumen, "Facturas": facturas, "Detalle": pd.DataFrame(det),
            "Revisar": pd.DataFrame(rev, columns=["Archivo", "Proveedor", "Número", "Problema"])}


def a_excel(hojas: dict[str, pd.DataFrame], destino: Path) -> None:
    with pd.ExcelWriter(destino, engine="openpyxl") as xl:
        for nombre, df in hojas.items():
            df.to_excel(xl, sheet_name=nombre, index=False)
            ws = xl.sheets[nombre]
            for celda in ws[1]:
                celda.fill, celda.font = AZUL, Font(bold=True, color="FFFFFF")
                celda.alignment = Alignment(horizontal="center")
            for j, col in enumerate(df.columns, start=1):
                ancho = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
                ws.column_dimensions[get_column_letter(j)].width = min(ancho, 60)
                if col in ("Subtotal", "IVA", "Total", "Valor unitario", "Valor total"):
                    for c in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        c[0].number_format = MONEDA
            if nombre == "Facturas":
                estado = list(df.columns).index("Estado") + 1
                for fila in ws.iter_rows(min_row=2):
                    if fila[estado - 1].value == "REVISAR":
                        for c in fila:
                            c.fill = ROJO
            ws.freeze_panes = "A2"


def reporte(hojas: dict[str, pd.DataFrame]) -> str:
    f, r, rev = hojas["Facturas"], hojas["Resumen"], hojas["Revisar"]
    ok = f[f.Estado == "OK"]
    pesos = lambda v: f"${v:,.0f}".replace(",", ".")
    L = [f"REPORTE DE FACTURAS — {datetime.now():%d/%m/%Y %H:%M}", "",
         f"Procesadas: {len(f)} · OK: {len(ok)} · Para revisar: {len(f) - len(ok)}",
         f"Total aprobado: {pesos(ok.Total.sum())} (IVA {pesos(ok.IVA.sum())})", "", "Por proveedor:"]
    for prov, tot in ok.groupby("Proveedor").Total.sum().sort_values(ascending=False).items():
        L.append(f"  · {prov}: {pesos(tot)}")
    if not rev.empty:
        L += ["", "Requieren revisión humana:"]
        L += [f"  ⚠ {x.Archivo}: {x.Problema}" for x in rev.itertuples()]
    return "\n".join(L)
